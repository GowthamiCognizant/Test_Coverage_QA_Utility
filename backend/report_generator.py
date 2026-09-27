import json
import subprocess
import sys
from pathlib import Path

_JS_SCRIPT = Path(__file__).parent / "generate_report.js"


def generate_word_report(analysis: dict, output_dir: Path) -> str:
    """
    Generates a rich Word (.docx) coverage report using docx-js via Node.js.
    Matches the Login Module Regression Coverage Report format:
      - Stats bar (Total / Covered / Uncovered / New Generated)
      - Coverage progress bar
      - Story-by-story table with Status (COVERED/PARTIAL/NEW) and C/R/N
      - Uncovered flows detail table
      - Feature file coverage summary table
    """
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    report_path = output_dir / "coverage_report.docx"

    # Pass analysis JSON as a command-line argument to the Node script
    analysis_json = json.dumps(analysis)

    try:
        result = subprocess.run(
            ["node", str(_JS_SCRIPT), analysis_json, str(report_path)],
            capture_output=True,
            text=True,
            timeout=60,
        )
        if result.returncode != 0:
            raise RuntimeError(
                f"Node.js report generation failed:\n{result.stderr}\n{result.stdout}"
            )
        output = result.stdout.strip()
        if not output.startswith("OK:"):
            raise RuntimeError(f"Unexpected output from report generator: {output}")

        return str(report_path)

    except FileNotFoundError:
        # Node.js not found — fall back to simple plain Word doc
        return _fallback_docx(analysis, report_path)
    except subprocess.TimeoutExpired:
        raise RuntimeError("Report generation timed out after 60 seconds")


def _fallback_docx(analysis: dict, report_path: Path) -> str:
    """Plain Word fallback if Node.js is not available."""
    try:
        from docx import Document
        from docx.shared import Pt, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        import datetime

        doc = Document()
        s = analysis.get("summary", {})

        title = doc.add_heading("Test Coverage Report", 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER

        doc.add_paragraph(
            f"Generated: {datetime.datetime.now().strftime('%d %B %Y %H:%M')}"
        )
        doc.add_paragraph("")
        doc.add_heading("Coverage Summary", 1)

        table = doc.add_table(rows=2, cols=4)
        table.style = "Table Grid"
        headers = ["Total TCs", "Covered", "Uncovered", "Coverage %"]
        values  = [
            str(s.get("total", 0)),
            str(s.get("covered", 0)),
            str(s.get("uncovered", 0)),
            f"{s.get('coverage_pct', 0)}%",
        ]
        for i, (h, v) in enumerate(zip(headers, values)):
            table.rows[0].cells[i].text = h
            table.rows[0].cells[i].paragraphs[0].runs[0].bold = True
            table.rows[1].cells[i].text = v

        doc.add_paragraph("")
        doc.add_heading("Uncovered Flows", 1)
        for flow, tcs in analysis.get("uncovered_flows", {}).items():
            doc.add_heading(flow, 2)
            t = doc.add_table(rows=1, cols=2)
            t.style = "Table Grid"
            t.rows[0].cells[0].text = "TC ID"
            t.rows[0].cells[1].text = "Description"
            for cell in t.rows[0].cells:
                cell.paragraphs[0].runs[0].bold = True
            for tc in tcs:
                r = t.add_row().cells
                r[0].text = tc.get("key", "")
                r[1].text = tc.get("summary", "")
            doc.add_paragraph("")

        doc.save(str(report_path))
        return str(report_path)

    except Exception as e:
        raise RuntimeError(f"Both Node.js and python-docx fallback failed: {e}")